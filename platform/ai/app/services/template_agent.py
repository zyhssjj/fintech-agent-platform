from __future__ import annotations

import re
from dataclasses import dataclass

from docx import Document

from app.clients.openai_compatible import OpenAICompatibleClient
from app.schemas import (
    TemplateChatRequest,
    TemplateDraftResponse,
    TemplateField,
    TemplateImportRequest,
    TemplateRenderAssistRequest,
)


CLIENT = OpenAICompatibleClient()
MAX_FOLLOW_UP_ROUNDS = 3

QUESTION_BANK: dict[str, str] = {
    "audience": "这份模板主要面向谁？例如 LP、被投企业，还是合作方？",
    "scenario": "这次沟通的具体场景是什么？例如季度净值通知、投后事项通知或合作函件。",
    "deliverable": "本次最终产物偏向哪一种？邮件正文、正式通知函，还是两者都要？",
    "required_fields": "这份模板每次生成时，哪些字段必须由业务人员手工补齐？",
    "tone": "语气希望更正式严谨，还是更偏服务型、解释型？",
}


@dataclass
class TemplateContext:
    audience: str = ""
    scenario: str = ""
    deliverable: str = ""
    tone: str = ""
    required_fields: list[str] | None = None


def _normalize_placeholders(text: str) -> str:
    return re.sub(r"(?<!\{)\{([a-zA-Z0-9_]+)\}(?!\})", r"{{\1}}", text)


def _plain_text_to_html(text: str) -> str:
    normalized = _normalize_placeholders(text.strip())
    if not normalized:
        return ""
    if any(tag in normalized for tag in ("<p", "<div", "<br")):
        return normalized
    paragraphs: list[str] = []
    for block in re.split(r"\n\s*\n", normalized):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        paragraphs.append(f"<p>{'<br/>'.join(lines)}</p>")
    return "".join(paragraphs)


def _strip_template_control_text(text: str) -> str:
    return (
        re.sub(r"\s*追问轮次[:：]\s*\d+[\s\S]*$", "", text, flags=re.IGNORECASE)
        .replace("请确认右侧草稿", "")
        .replace("如无问题可直接保存", "")
        .replace("草稿确认阶段", "")
        .strip()
    )


def _is_instructional_text(text: str) -> bool:
    normalized = text.strip()
    if not normalized:
        return False
    markers = (
        "请回复确认",
        "回复“确认”",
        "回复\"确认\"",
        "确认开始优化",
        "如需我继续",
        "如需我直接继续",
        "请直接提供现有草稿内容",
        "请确认开始优化",
    )
    return any(marker in normalized for marker in markers)


def _is_valid_template_fragment(text: str) -> bool:
    normalized = _strip_template_control_text(text)
    if not normalized or _is_instructional_text(normalized):
        return False
    has_placeholder = bool(re.search(r"\{\{?[a-zA-Z0-9_]+\}?\}", normalized))
    has_html = any(tag in normalized.lower() for tag in ("<p", "<div", "<br"))
    long_enough = len(re.sub(r"\s+", "", normalized)) >= 30
    return has_placeholder or has_html or long_enough


def _extract_summary_section(text: str, labels: list[str]) -> str:
    lines = text.replace("\r\n", "\n").splitlines()
    escaped = [re.escape(label) for label in labels]
    start_pattern = re.compile(rf"^(?:{'|'.join(escaped)})\s*[:：]?\s*(.*)$", re.IGNORECASE)
    stop_pattern = re.compile(
        r"^(\u5efa\u8bae\u5b57\u6bb5|\u5f85\u786e\u8ba4\u95ee\u9898|\u8bf7\u8865\u5145|\u8ffd\u95ee\u8f6e\u6b21|subject template|html template)\s*[:：]?",
        re.IGNORECASE,
    )
    content: list[str] = []
    started = False

    for line in lines:
        stripped = line.strip()
        if not started:
            matched = start_pattern.match(stripped)
            if not matched:
                continue
            started = True
            if matched.group(1).strip():
                content.append(matched.group(1).strip())
            continue
        if stripped and stop_pattern.match(stripped):
            break
        content.append(line)

    return _strip_template_control_text("\n".join(content).strip())


def _extract_structured_draft(text: str) -> tuple[str, str, list[str]]:
    normalized = text.replace("\r\n", "\n").strip()
    extracted_subject = _extract_summary_section(
        normalized,
        ["\u4e3b\u9898\u6a21\u677f", "subject template", "subject"],
    )
    extracted_html = _extract_summary_section(
        normalized,
        ["HTML \u6a21\u677f", "HTML\u6a21\u677f", "html template", "html"],
    )
    if extracted_subject or extracted_html:
        variable_names: set[str] = set()
        for source in (extracted_subject, extracted_html, normalized):
            for matched in re.findall(r"\{+([a-zA-Z0-9_]+)\}+", source):
                variable_names.add(matched.strip())
            for matched in re.findall(r"\b([a-zA-Z]+_[a-zA-Z0-9_]+)\b", source):
                variable_names.add(matched.strip())
        return (
            _normalize_placeholders(extracted_subject) if _is_valid_template_fragment(extracted_subject) else "",
            _plain_text_to_html(_strip_template_control_text(extracted_html)) if _is_valid_template_fragment(extracted_html) else "",
            sorted(variable_names),
        )

    raw_subject_lines: list[str] = []
    raw_html_lines: list[str] = []
    section = ""
    for line in normalized.splitlines():
        stripped = line.strip()
        if not stripped:
            if section == "html":
                raw_html_lines.append("")
            elif section == "subject":
                raw_subject_lines.append("")
            continue
        if "HTML" in stripped and ("模板" in stripped or "Template" in stripped):
            section = "html"
            after = re.split(r"[:：]", stripped, maxsplit=1)
            if len(after) == 2 and after[1].strip():
                raw_html_lines.append(after[1].strip())
            continue
        if ("主题" in stripped and "模板" in stripped) or stripped.lower().startswith("subject template"):
            section = "subject"
            after = re.split(r"[:：]", stripped, maxsplit=1)
            if len(after) == 2 and after[1].strip():
                raw_subject_lines.append(after[1].strip())
            continue
        if any(marker in stripped for marker in ("建议字段", "待确认问题", "追问轮次")):
            section = ""
            continue
        if section == "subject":
            raw_subject_lines.append(stripped)
        elif section == "html":
            raw_html_lines.append(stripped)
    raw_subject = "\n".join(raw_subject_lines).strip()
    raw_html = "\n".join(raw_html_lines).strip()
    variable_names: set[str] = set()
    for source in (raw_subject, raw_html, normalized):
        for matched in re.findall(r"\{+([a-zA-Z0-9_]+)\}+", source):
            variable_names.add(matched.strip())
        for matched in re.findall(r"\b([a-zA-Z]+_[a-zA-Z0-9_]+)\b", source):
            variable_names.add(matched.strip())
    return _normalize_placeholders(raw_subject), _plain_text_to_html(raw_html), sorted(variable_names)


def _merge_structured_response(
    response: TemplateDraftResponse,
    ai_text: str,
    fallback_fields: list[TemplateField],
) -> TemplateDraftResponse:
    subject_template, html_template, variable_names = _extract_structured_draft(ai_text)
    if not subject_template and not html_template:
        subject_template, html_template, variable_names = _extract_structured_draft(response.summary or "")
    field_map = {field.variable_name: field for field in (response.suggested_fields or fallback_fields)}
    for name in variable_names:
        if name not in field_map:
            field_map[name] = TemplateField(
                variable_name=name,
                variable_label=name,
                variable_type="date" if re.search(r"date|time|period", name, re.IGNORECASE) else "text",
            )
    return TemplateDraftResponse(
        summary=response.summary,
        follow_up_questions=response.follow_up_questions,
        suggested_fields=list(field_map.values()),
        subject_template=subject_template or response.subject_template,
        html_template=html_template or response.html_template,
    )


def _normalize_text(*parts: str) -> str:
    return "\n".join(part.strip() for part in parts if part and part.strip())


def _fallback_fields(scene: str, audience: str = "") -> list[TemplateField]:
    scene_text = f"{scene} {audience}".lower()
    fields = [
        TemplateField(variable_name="fund_name", variable_label="基金名称"),
        TemplateField(variable_name="notice_date", variable_label="通知日期", variable_type="date"),
    ]
    if "lp" in scene_text:
        fields.insert(1, TemplateField(variable_name="lp_name", variable_label="LP 名称"))
    elif "合作" in scene or "partner" in scene_text:
        fields.insert(1, TemplateField(variable_name="partner_name", variable_label="合作方名称"))
    elif "被投" in scene or "portfolio" in scene_text:
        fields.insert(1, TemplateField(variable_name="portfolio_name", variable_label="被投企业名称"))
    else:
        fields.insert(1, TemplateField(variable_name="recipient_name", variable_label="收件对象名称"))
    return fields


def _extract_context(text: str, known_fields: list[dict] | None = None) -> TemplateContext:
    lower = text.lower()
    context = TemplateContext(required_fields=[])

    if "lp" in lower:
        context.audience = "LP"
    elif "合作方" in text or "partner" in lower:
        context.audience = "合作方"
    elif "被投" in text or "portfolio" in lower:
        context.audience = "被投企业"

    if "季度" in text and "通知" in text:
        context.scenario = "季度通知"
    elif "净值" in text:
        context.scenario = "净值通知"
    elif "投后" in text:
        context.scenario = "投后管理通知"
    elif "合作" in text:
        context.scenario = "合作沟通函件"
    elif "通知" in text:
        context.scenario = "正式通知"

    if "docx" in lower or "函件" in text or "word" in lower:
        context.deliverable = "正式通知函"
    if "邮件" in text or "email" in lower:
        context.deliverable = "邮件正文" if not context.deliverable else "邮件正文 + 正式通知函"

    if "正式" in text or "严谨" in text or "专业" in text:
        context.tone = "正式严谨"
    elif "温和" in text or "服务" in text or "解释" in text:
        context.tone = "服务解释型"

    field_names = set()
    for matched in re.findall(r"{{\s*([a-zA-Z0-9_]+)\s*}}", text):
        field_names.add(matched.strip())
    for matched in re.findall(r"\b([a-zA-Z]+_[a-zA-Z0-9_]+)\b", text):
        field_names.add(matched.strip())
    if known_fields:
        for field in known_fields:
            name = str(field.get("variable_name", "")).strip()
            if name:
                field_names.add(name)
    context.required_fields = sorted(field_names)
    return context


def _guess_round(current_summary: str) -> int:
    matched = re.search(r"追问轮次[:：]\s*(\d+)", current_summary)
    if not matched:
        return 0
    try:
        return int(matched.group(1))
    except ValueError:
        return 0


def _missing_topics(context: TemplateContext) -> list[str]:
    missing: list[str] = []
    if not context.audience:
        missing.append("audience")
    if not context.scenario:
        missing.append("scenario")
    if not context.deliverable:
        missing.append("deliverable")
    if not context.required_fields:
        missing.append("required_fields")
    if not context.tone:
        missing.append("tone")
    return missing


def _subject_template(context: TemplateContext) -> str:
    if context.audience == "LP":
        return "关于{{fund_name}}的正式通知"
    if context.audience == "合作方":
        return "关于{{partner_name}}合作事项的正式通知"
    if context.audience == "被投企业":
        return "关于{{portfolio_name}}相关事项的正式通知"
    return "关于{{fund_name}}相关事项的正式通知"


def _html_template(context: TemplateContext, fields: list[TemplateField]) -> str:
    recipient_field = "recipient_name"
    for field in fields:
        if field.variable_name in {"lp_name", "partner_name", "portfolio_name", "recipient_name"}:
            recipient_field = field.variable_name
            break
    intro = "现将相关事项通知如下。"
    if context.scenario:
        intro = f"现将{context.scenario}相关事项通知如下。"
    return (
        f"<p>尊敬的{{{{{recipient_field}}}}}：</p>"
        f"<p>{intro}</p>"
        "<p>基金名称：{{fund_name}}</p>"
        "<p>通知日期：{{notice_date}}</p>"
        "<p>如有疑问，请及时与我们联系。</p>"
    )


def _summary_text(context: TemplateContext, round_no: int, force_confirm: bool) -> str:
    parts = ["已更新模板草稿。"]
    if context.audience:
        parts.append(f"面向对象：{context.audience}。")
    if context.scenario:
        parts.append(f"模板场景：{context.scenario}。")
    if context.deliverable:
        parts.append(f"产物类型：{context.deliverable}。")
    if context.required_fields:
        parts.append(f"建议字段：{', '.join(context.required_fields)}。")
    if context.tone:
        parts.append(f"语气风格：{context.tone}。")
    if force_confirm:
        parts.append("请确认右侧草稿；如无问题，可直接手改后保存。")
    else:
        parts.append("我会继续围绕缺失信息做有限追问。")
    parts.append(f"追问轮次：{round_no}")
    return "".join(parts)


def _merge_ai_summary(ai_text: str, fallback_summary: str, round_no: int) -> str:
    if ai_text and ai_text.strip():
        cleaned = _strip_template_control_text(ai_text.strip())
        base = fallback_summary if _is_instructional_text(cleaned) else cleaned
    else:
        base = fallback_summary
    if f"追问轮次：{round_no}" in base:
        return base
    return f"{base} 追问轮次：{round_no}"


def _question_list(current_summary: str, missing: list[str], round_no: int) -> list[str]:
    asked = {line.strip() for line in current_summary.splitlines() if line.strip()}
    questions: list[str] = []
    for topic in missing:
        question = QUESTION_BANK[topic]
        if question in asked:
            continue
        questions.append(question)
        if len(questions) >= 2:
            break
    if round_no >= MAX_FOLLOW_UP_ROUNDS:
        return []
    return questions


async def draft_template_from_chat(payload: TemplateChatRequest) -> TemplateDraftResponse:
    text = _normalize_text(payload.scene, payload.goal, payload.current_summary, payload.latest_user_message)
    context = _extract_context(text, payload.known_fields)
    round_no = min(_guess_round(payload.current_summary) + 1, MAX_FOLLOW_UP_ROUNDS)
    missing = _missing_topics(context)
    follow_ups = _question_list(payload.current_summary, missing, round_no)
    force_confirm = round_no >= MAX_FOLLOW_UP_ROUNDS or (round_no >= 2 and len(missing) <= 1) or (round_no >= 2 and not follow_ups)
    fields = _fallback_fields(payload.scene or payload.goal, context.audience)

    if context.required_fields:
        field_names = {field.variable_name for field in fields}
        for name in context.required_fields:
            if name not in field_names:
                fields.append(TemplateField(variable_name=name, variable_label=name.replace("_", " ").title()))

    fallback = TemplateDraftResponse(
        summary=_summary_text(context, round_no, force_confirm),
        follow_up_questions=[] if force_confirm else follow_ups,
        suggested_fields=fields,
        subject_template=_subject_template(context),
        html_template=_html_template(context, fields),
    )
    if not CLIENT.is_configured():
        return fallback

    system_prompt = (
        "你是金融科技通讯管理模板助手。"
        "你只能围绕面向对象、模板场景、产物类型、必填字段、语气风格五类信息追问。"
        "同一个问题不能重复问，最多连续追问三轮。"
        "如果最小信息已满足，或者已经到第三轮，就直接要求用户确认草稿。"
        "输出简洁中文，不输出思考过程，不输出 Markdown。"
    )
    user_prompt = (
        f"场景：{payload.scene}\n"
        f"模板目标：{payload.goal}\n"
        f"当前摘要：{payload.current_summary}\n"
        f"用户最新输入：{payload.latest_user_message}\n"
        f"已知字段：{payload.known_fields}\n"
        f"当前回合：{round_no}\n"
        "请基于这些信息优化模板摘要、主题模板和 HTML 模板。"
    )
    try:
        text = await CLIENT.chat(system_prompt, user_prompt)
    except Exception:
        return fallback
    fallback.summary = _merge_ai_summary(text or "", fallback.summary, round_no)
    return _merge_structured_response(fallback, text or "", fields)


def _extract_candidate_fields(text: str) -> list[TemplateField]:
    fields: list[TemplateField] = []
    seen: set[str] = set()
    mapping = [
        ("基金", "fund_name", "基金名称"),
        ("LP", "lp_name", "LP 名称"),
        ("合作方", "partner_name", "合作方名称"),
        ("被投", "portfolio_name", "被投企业名称"),
        ("日期", "notice_date", "通知日期", "date"),
    ]
    for item in mapping:
        keyword, variable_name, label, *rest = item
        if keyword in text and variable_name not in seen:
            seen.add(variable_name)
            fields.append(
                TemplateField(
                    variable_name=variable_name,
                    variable_label=label,
                    variable_type=rest[0] if rest else "text",
                )
            )
    if not fields:
        fields = _fallback_fields(text)
    return fields


async def draft_template_from_import(payload: TemplateImportRequest) -> TemplateDraftResponse:
    text = payload.plain_text.strip() or "\n".join(item for item in payload.paragraphs if item.strip())
    context = _extract_context(_normalize_text(payload.scene, text))
    fields = _extract_candidate_fields(text)
    summary = f"已从 {payload.file_name} 识别出 {len(fields)} 个候选字段，请确认草稿并补充缺失信息。"
    follow_ups = []
    missing = _missing_topics(context)
    if missing:
        follow_ups = _question_list("", missing, 1)

    response = TemplateDraftResponse(
        summary=summary,
        follow_up_questions=follow_ups,
        suggested_fields=fields,
        subject_template=_subject_template(context),
        html_template=_html_template(context, fields),
    )
    if not CLIENT.is_configured():
        return response

    system_prompt = (
        "你是金融科技通讯管理模板解析助手。"
        "请从导入的 Word / 历史函件中识别文档用途、候选字段、结构判断和缺失信息。"
        "不要输出思考过程，不要输出 Markdown。"
    )
    user_prompt = (
        f"文件名：{payload.file_name}\n"
        f"模板场景：{payload.scene}\n"
        f"文档文本：\n{text[:6000]}\n"
        "请给出简短摘要，并优化主题模板与 HTML 模板。"
    )
    try:
        ai_summary = await CLIENT.chat(system_prompt, user_prompt)
    except Exception:
        return response
    if ai_summary and ai_summary.strip():
        response.summary = ai_summary.strip()
    return _merge_structured_response(response, ai_summary or "", fields)


async def assist_render(payload: TemplateRenderAssistRequest) -> TemplateDraftResponse:
    recipient = payload.recipient_type or "外部收件人"
    response = TemplateDraftResponse(
        summary="已生成一版正式沟通正文建议。",
        follow_up_questions=[],
        suggested_fields=[],
        subject_template=payload.template_name or "正式通知",
        html_template=(
            f"<p>尊敬的{recipient}：</p>"
            "<p>现将相关事项说明如下。</p>"
            "<p>请结合附件或正式函件内容查收。</p>"
        ),
    )
    if not CLIENT.is_configured():
        return response

    system_prompt = (
        "你是金融科技对外沟通文案助手。"
        "请输出正式、专业、清晰的邮件主题和邮件正文建议。"
        "正文只能使用已知变量中的真实值；未知发件人姓名、职位、单位或联系方式应直接省略，"
        "绝不能输出 [发件人姓名]、[联系方式]、{发件人姓名}、{{sender_name}} 等未填充占位符。"
        "不要输出思考过程，不要输出 Markdown。"
    )
    user_prompt = (
        f"模板名称：{payload.template_name}\n"
        f"模板场景：{payload.template_scene}\n"
        f"收件对象类型：{payload.recipient_type}\n"
        f"业务目标：{payload.business_goal}\n"
        f"已知变量：{payload.known_values}\n"
        "请优化一版主题和正文。"
    )
    try:
        text = await CLIENT.chat(system_prompt, user_prompt)
    except Exception:
        return response
    if text and text.strip():
        subject, html = _extract_render_assist_draft(text)
        response.summary = "\u5df2\u751f\u6210\u6b63\u5f0f\u6c9f\u901a\u6b63\u6587\u5efa\u8bae\uff0c\u53ef\u5728\u4e0b\u65b9\u7f16\u8f91\u540e\u9884\u89c8\u6216\u521b\u5efa\u51fd\u4ef6\u3002"
        response.subject_template = subject or response.subject_template
        response.html_template = html or response.html_template
    return response


def _extract_render_assist_draft(text: str) -> tuple[str, str]:
    """Turn the model's human-readable mail suggestion into the fields used by the UI."""
    lines = [line.strip() for line in text.replace("\r\n", "\n").split("\n")]
    subject = ""
    body_lines: list[str] = []
    seen_subject = False
    subject_pattern = re.compile(r"^\s*(?:\u90ae\u4ef6)?\u4e3b\u9898(?:\u5efa\u8bae)?\s*[:\uff1a]\s*(.+)$", re.IGNORECASE)
    body_pattern = re.compile(r"^\s*(?:\u90ae\u4ef6)?\u6b63\u6587(?:\u5efa\u8bae|\u9884\u89c8)?\s*[:\uff1a]?\s*(.*)$", re.IGNORECASE)

    for line in lines:
        subject_match = subject_pattern.match(line)
        if subject_match and not subject:
            subject = subject_match.group(1).strip()
            seen_subject = True
            continue
        body_match = body_pattern.match(line)
        if body_match:
            seen_subject = True
            if body_match.group(1).strip():
                body_lines.append(body_match.group(1).strip())
            continue
        if seen_subject:
            body_lines.append(line)

    body = "\n".join(body_lines).strip()
    return _normalize_placeholders(subject), _plain_text_to_html(_strip_template_control_text(body)) if body else ""


def extract_docx_text(file_path: str) -> tuple[str, list[str]]:
    document = Document(file_path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs), paragraphs
