"""Regression coverage for chapter-aware long-document retrieval."""

from __future__ import annotations

import unittest
from unittest.mock import patch

from scripts import retrieval
from scripts.ingest_finance import DocumentReadError, read_document, structured_chunk, validate_structured_plans


def record(record_id: str, chunk_index: int, section_title: str, content: str, company_code: str = "beijing_hangke") -> dict:
    return {
        "id": record_id,
        "score": 1.0,
        "source": {
            "doc_id": "doc-architecture",
            "chunk_id": record_id,
            "chunk_index": chunk_index,
            "title": "系统架构设计",
            "chunk_title": section_title,
            "section_title": section_title,
            "table_header": "序号 | 名称 | 说明",
            "content": content,
            "company_code": company_code,
            "domain_path": "公司层面/公司运营类",
        },
    }


class FakeEmbedding:
    def embed_text(self, _query: str) -> list[float]:
        return [0.1, 0.2]


class FakeElasticsearch:
    """Small ES double that preserves the filters used by retrieval.py."""

    def __init__(self, rows: list[dict]):
        self.rows = rows

    def search(self, *, query=None, knn=None, sort=None, **_kwargs):
        filters = []
        if query:
            filters = query.get("bool", {}).get("filter", [])
        elif knn:
            filters = knn.get("filter", [])
        rows = self.rows
        for item in filters:
            term = item.get("term", {})
            for key, value in term.items():
                source_key = key.removesuffix(".keyword")
                rows = [row for row in rows if row["source"].get(source_key) == value]
        if sort:
            rows = sorted(rows, key=lambda row: row["source"].get("chunk_index", 0))
        return {"hits": {"hits": [{"_id": row["id"], "_source": row["source"], "_score": row["score"]} for row in rows]}}


class RetrievalRegressionTests(unittest.TestCase):
    def setUp(self):
        self.rows = [
            record("system-0", 0, "1.1 系统定位", "【章节】1.1 系统定位\n第一段定位说明。"),
            record("system-1", 1, "1.1 系统定位", "【章节】1.1 系统定位\n第二段定位说明。"),
            record("firewall-0", 2, "2. 防火墙（共2项）", "【章节】2. 防火墙（共2项）\n【表头】序号 | 名称 | 说明\n【表格数据】\n1 | FW-A | 第一项"),
            record("firewall-1", 3, "2. 防火墙（共2项）", "【章节】2. 防火墙（共2项）\n【表头】序号 | 名称 | 说明\n【表格数据】\n2 | FW-B | 第二项"),
            record("other-company", 4, "2. 防火墙（共1项）", "【章节】2. 防火墙（共1项）\n【表格数据】\n1 | OTHER | 不可跨公司", "beijing_motian"),
        ]
        self.es = FakeElasticsearch(self.rows)
        self.emb = FakeEmbedding()

    def retrieve(self, query: str, company_code: str = "beijing_hangke"):
        with patch.object(retrieval, "_rerank", side_effect=lambda _query, candidates: (candidates, False)):
            return retrieval.retrieve(self.es, self.emb, "finance", query, final_k=10, filters={"company_code": company_code})

    def test_exact_list_subject_expands_only_its_section_in_order(self):
        hits, meta = self.retrieve("防火墙有哪些")
        expanded = [item for item in hits if item.get("section_expanded")]
        self.assertTrue(meta["list_query"])
        self.assertEqual(meta["query_subject"], "防火墙")
        self.assertEqual([item["source"]["chunk_index"] for item in expanded], [2, 3])
        self.assertTrue(all(item["source"]["section_title"].startswith("2. 防火墙") for item in expanded))

    def test_normal_question_anchors_matching_section(self):
        hits, meta = self.retrieve("系统定位是什么")
        expanded = [item for item in hits if item.get("section_expanded")]
        self.assertFalse(meta["list_query"])
        self.assertEqual([item["source"]["chunk_index"] for item in expanded], [0, 1])

    def test_missing_list_does_not_expand_semantically_related_section(self):
        hits, meta = self.retrieve("测试工具有哪些")
        self.assertTrue(meta["list_query"])
        self.assertFalse(meta["section_anchor_found"])
        self.assertEqual(meta["section_expanded_count"], 0)
        self.assertFalse(any(item.get("section_expanded") for item in hits))

    def test_company_filter_prevents_cross_company_results(self):
        hits, _ = self.retrieve("防火墙有哪些", "beijing_motian")
        self.assertTrue(hits)
        self.assertTrue(all(item["source"]["company_code"] == "beijing_motian" for item in hits))

    def test_structured_chunks_repeat_heading_and_table_header_for_each_row_group(self):
        plans = structured_chunk(
            "# 防火墙\n| 序号 | 名称 | 说明 |\n| 1 | FW-A | 第一项 |\n| 2 | FW-B | 第二项 |",
            size=25,
            overlap=0,
        )
        table_plans = [plan for plan in plans if plan["is_structured"]]
        self.assertGreaterEqual(len(table_plans), 2)
        self.assertTrue(all("【章节】防火墙" in plan["content"] for plan in table_plans))
        self.assertTrue(all("【表格】第1个数据表" in plan["content"] for plan in table_plans))
        self.assertTrue(all("【表头】| 序号 | 名称 | 说明 |" in plan["content"] for plan in table_plans))
        self.assertTrue(all(plan["table_index"] == 1 for plan in table_plans))
        self.assertEqual([(plan["table_row_start"], plan["table_row_end"]) for plan in table_plans], [(1, 1), (2, 2)])
        self.assertEqual(validate_structured_plans(table_plans)["table_row_count"], 2)

    def test_numbered_table_rows_are_not_mistaken_for_section_headings(self):
        plans = structured_chunk(
            "# 防火墙\n| 序号 | 名称 | 说明 |\n1 | FW-A | 第一项 |\n2 | FW-B | 第二项 |",
            size=30,
            overlap=0,
        )
        table_plans = [plan for plan in plans if plan["is_structured"]]
        self.assertEqual(len(table_plans), 2)
        self.assertTrue(all(plan["section_title"] == "防火墙" for plan in table_plans))
        self.assertIn("1 | FW-A | 第一项 |", table_plans[0]["content"])
        self.assertIn("2 | FW-B | 第二项 |", table_plans[1]["content"])

    def test_numbered_dot_heading_starts_a_new_section(self):
        plans = structured_chunk("1. 防火墙\n这是用于验证编号标题识别的足够长说明内容，不能被最小长度过滤。", size=100, overlap=0)
        self.assertEqual(plans[0]["section_title"], "防火墙")

    def test_docx_reader_preserves_heading_before_its_table(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path
        from docx import Document

        with TemporaryDirectory() as directory:
            path = Path(directory) / "ordered.docx"
            document = Document()
            document.add_paragraph("1. 防火墙")
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "序号", "名称"
            table.rows[1].cells[0].text, table.rows[1].cells[1].text = "1", "FW-A"
            document.add_paragraph("2. 入侵检测")
            document.save(path)
            text, parser = read_document(path)
        self.assertEqual(parser, "docx")
        self.assertLess(text.index("1. 防火墙"), text.index("1 | FW-A"))
        self.assertLess(text.index("1 | FW-A"), text.index("2. 入侵检测"))
        self.assertIn("【表格】第1个数据表", text)

    def test_integrity_rejects_declared_count_that_does_not_match_rows(self):
        plans = structured_chunk(
            "# 防火墙（共3项）\n| 序号 | 名称 |\n1 | FW-A |\n2 | FW-B |",
            size=100,
            overlap=0,
        )
        with self.assertRaises(DocumentReadError):
            validate_structured_plans(plans)

    def test_multiple_tables_are_numbered_and_validated_independently(self):
        plans = structured_chunk(
            "【表格】第1个数据表\n甲 | 乙\n1 | A\n"
            "【表格】第2个数据表\n丙 | 丁\n1 | B\n2 | C",
            size=100,
            overlap=0,
        )
        tables = [plan for plan in plans if plan["is_structured"]]
        self.assertEqual([plan["table_index"] for plan in tables], [1, 2])
        self.assertEqual(validate_structured_plans(plans), {"table_count": 2, "table_row_count": 3})

    def test_section_expansion_is_not_limited_to_a_fixed_small_chunk_count(self):
        self.assertEqual(retrieval.FINANCE_PROFILE["section_expand_limit"], 10_000)


if __name__ == "__main__":
    unittest.main()
